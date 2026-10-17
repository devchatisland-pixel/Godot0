"""Orphan models audit: every model placed only once or twice on the map, sorted by zone.

  python tools/build_orphan_audit.py   (after tools/dump_building_models.ps1, render_zone_models.ps1
                                        and dump_fog_and_vehicles.ps1; tools/audit_orphans.ps1 runs all)

Reads docs/building_models.tsv (buildings of the main island and the islets) and
docs/fog_and_vehicles.tsv (fog island, vehicles, boats and props of the whole map), counts each model
over the whole map and writes docs/audit orphelin models.docx: one section per zone, a table
Picture | Model | Where | Count (Count = times placed on the whole map, Where = per zone).
"""
import csv, os, re
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt

DOCS = "docs"
OUT = os.path.join(DOCS, "audit orphelin models.docx")

ZONE_NAMES = {"NONE": "SEA"}

def safe(name):
    return re.sub(r"[^a-z0-9]", "_", name.lower())

def main():
    # model -> {"where": {zone: n}, "pic": path, "area": text}
    models = {}
    def add(model, zone, n, pic_dir, area):
        m = models.setdefault(model, {"where": defaultdict(int), "pic": os.path.join(DOCS, pic_dir, safe(model) + ".png"), "area": area})
        m["where"][ZONE_NAMES.get(zone, zone)] += n
    with open(os.path.join(DOCS, "building_models.tsv"), encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            if r["model"].startswith("(no model"): continue
            add(r["model"], r["zone"], int(r["count"]), "building_model_images", "Main island and islets")
    with open(os.path.join(DOCS, "fog_and_vehicles.tsv"), encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            area = "Fog island" if r["area"] == "FOG" else "Vehicles, boats and props"
            zone = ("FOG: " if r["area"] == "FOG" else "") + ZONE_NAMES.get(r["zone"], r["zone"])
            m = models.setdefault(r["label"], {"where": defaultdict(int), "pic": os.path.join(DOCS, "fog_vehicle_images", safe(r["key"]) + ".png"), "area": area})
            m["where"][zone] += int(r["count"])
    orphans = {k: v for k, v in models.items() if sum(v["where"].values()) <= 2}
    by_zone = defaultdict(list)
    for name, m in orphans.items():
        for zone in m["where"]:
            by_zone[zone].append(name)
    once = sum(1 for m in orphans.values() if sum(m["where"].values()) == 1)

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
    doc.add_heading("Orphan models audit", 0)
    doc.add_paragraph("Models that appear only once or twice on the map (Chat City, its islets and the fog "
                      "island; the continent is not counted). Count = how many times the model is placed on the "
                      "whole map; Where = how many of them in each zone. The models are sorted by zone: a model "
                      "placed in two zones is listed under both.")
    doc.add_heading("Summary", 1)
    doc.add_paragraph("%d models are used only once or twice: %d used once, %d used twice, out of %d different "
                      "models placed." % (len(orphans), once, len(orphans) - once, len(models)))
    for zone in sorted(by_zone):
        doc.add_paragraph("%s: %d models" % (zone, len(by_zone[zone])), style="List Bullet")
    for zone in sorted(by_zone):
        names = sorted(by_zone[zone], key=lambda n: (-sum(orphans[n]["where"].values()), n.lower()))
        doc.add_heading("%s (%d)" % (zone, len(names)), 1)
        t = doc.add_table(rows=1, cols=4)
        t.style = "Table Grid"
        for c, h in zip(t.rows[0].cells, ["Picture", "Model", "Where", "Count"]):
            c.text = h
            c.paragraphs[0].runs[0].bold = True
        for n in sorted(names, key=str.lower):
            m = orphans[n]
            row = t.add_row().cells
            if os.path.exists(m["pic"]):
                row[0].paragraphs[0].add_run().add_picture(m["pic"], width=Cm(2.4))
            row[1].text = n
            row[2].text = ", ".join("%s:%d" % (z, k) for z, k in sorted(m["where"].items()))
            row[3].text = str(sum(m["where"].values()))
        for r in t.rows:
            for c, w in zip(r.cells, [Cm(2.8), Cm(6.5), Cm(6), Cm(1.5)]):
                c.width = w
                for p in c.paragraphs:
                    for run in p.runs:
                        run.font.size = Pt(9)
    doc.save(OUT)
    missing = [n for n, m in orphans.items() if not os.path.exists(m["pic"])]
    print("orphans %d (%d once), zones %d, no picture: %d %s" % (len(orphans), once, len(by_zone), len(missing), missing[:10]))

if __name__ == "__main__":
    main()
