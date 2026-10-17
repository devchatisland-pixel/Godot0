"""Resources audit: every model placed on the map, sorted by the memory it takes (biggest first).

  python tools/build_resource_audit.py   (after tools/audit_orphans.ps1, which runs it too)

Reads docs/building_models.tsv and docs/fog_and_vehicles.tsv (where each model stands and how often)
and docs/used_models.tsv (tests/dump_used_models.gd: memory of each model = its mesh data and the
textures of its materials, in MB; loaded once whatever the number of buildings) and writes
docs/audit resources models.docx: Picture | Model | Zone | Buildings | Memory (MB).
"""
import csv, os, re
from collections import defaultdict
from docx import Document
from docx.shared import Cm, Pt

DOCS = "docs"
OUT = os.path.join(DOCS, "audit resources models.docx")
ZONE_NAMES = {"NONE": "SEA"}

def safe(name):
    return re.sub(r"[^a-z0-9]", "_", name.lower())

def main():
    mem = {}
    with open(os.path.join(DOCS, "used_models.tsv"), encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            mem[r["model"]] = (float(r["memory_mb"]), int(r["triangles"]))
    models = {}
    def add(label, key, zone, n, pic_dir):
        m = models.setdefault(label, {"key": key, "where": defaultdict(int),
                                      "pic": os.path.join(DOCS, pic_dir, safe(key) + ".png")})
        m["where"][ZONE_NAMES.get(zone, zone)] += n
    with open(os.path.join(DOCS, "building_models.tsv"), encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            if not r["model"].startswith("(no model"):
                add(r["model"], r["model"], r["zone"], int(r["count"]), "building_model_images")
    with open(os.path.join(DOCS, "fog_and_vehicles.tsv"), encoding="utf-8") as f:
        for r in csv.DictReader(f, delimiter="\t"):
            zone = ("FOG: " if r["area"] == "FOG" else "") + ZONE_NAMES.get(r["zone"], r["zone"])
            add(r["label"], r["key"], zone, int(r["count"]), "fog_vehicle_images")
    rows = []
    for label, m in models.items():
        mb, tris = mem.get(m["key"], (None, 0))
        rows.append((mb if mb is not None else -1.0, label, m))
    rows.sort(key=lambda r: (-r[0], r[1].lower()))
    known = [r for r in rows if r[0] >= 0]
    total = sum(r[0] for r in known)

    doc = Document()
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(1.8)
    doc.add_heading("Resources audit of the models", 0)
    doc.add_paragraph("Every model placed on the map (Chat City, its islets, the fog island, the vehicles; the "
                      "continent is not counted), sorted by Memory: what the model takes once loaded, its mesh "
                      "data (vertices and triangles) plus the textures of its materials, in MB. A model is loaded "
                      "once however many buildings use it, so the biggest ones are where loading time and memory "
                      "go. Buildings = how many times it is placed; Zone = how many in each zone.")
    doc.add_heading("Summary", 1)
    doc.add_paragraph("%d models placed, %.1f MB in all. The 10 biggest take %.1f MB (%d%%)." % (
        len(rows), total, sum(r[0] for r in known[:10]), round(100 * sum(r[0] for r in known[:10]) / max(total, 0.001))))
    missing = [r[1] for r in rows if r[0] < 0]
    if missing:
        doc.add_paragraph("No memory figure for %d models (drawn outside the model library): %s." % (
            len(missing), ", ".join(missing)))
    doc.add_heading("Models by memory", 1)
    t = doc.add_table(rows=1, cols=5)
    t.style = "Table Grid"
    for c, h in zip(t.rows[0].cells, ["Picture", "Model", "Zone", "Buildings", "Memory (MB)"]):
        c.text = h
        c.paragraphs[0].runs[0].bold = True
    for mb, label, m in rows:
        row = t.add_row().cells
        if os.path.exists(m["pic"]):
            row[0].paragraphs[0].add_run().add_picture(m["pic"], width=Cm(2.2))
        row[1].text = label
        row[2].text = ", ".join("%s:%d" % (z, k) for z, k in sorted(m["where"].items()))
        row[3].text = str(sum(m["where"].values()))
        row[4].text = "%.2f" % mb if mb >= 0 else "-"
    for r in t.rows:
        for c, w in zip(r.cells, [Cm(2.6), Cm(5.6), Cm(5.2), Cm(1.8), Cm(2.0)]):
            c.width = w
            for p in c.paragraphs:
                for run in p.runs:
                    run.font.size = Pt(9)
    doc.save(OUT)
    print("resource audit: %d models, %.1f MB, no figure for %d, no picture for %d" % (
        len(rows), total, len(missing), sum(1 for r in rows if not os.path.exists(r[2]["pic"]))))

if __name__ == "__main__":
    main()
